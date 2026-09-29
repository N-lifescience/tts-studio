"""모델 없이 도는 테스트. 생성·받아쓰기는 가짜로 바꿔 끼운다."""

import os
import time
from pathlib import Path

import numpy as np
import pytest
import soundfile as sf

os.environ["STUDIO_TESTING"] = "1"

from server import check, config, engine, exporter, store, subtitles, worker  # noqa: E402
from server.textsplit import join_script, split_script  # noqa: E402

SCRIPT = """오늘은 이 공허의 약탈자를, 게임이 아니라 생물학의 눈으로 한 번 뜯어보겠습니다.
먼저 몸부터 보죠. 몸을 빠르게 움직이는 뒤쪽 다리

사마귀입니다."""


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PROJECTS_DIR", tmp_path / "projects")
    monkeypatch.setattr(config, "VOICES_DIR", tmp_path / "voices")
    monkeypatch.setattr(config, "EXPORT_ROOT", tmp_path / "icloud" / "TTS")
    (tmp_path / "icloud").mkdir()
    config.VOICES_DIR.mkdir()
    sr = config.SR
    tone = 0.3 * np.sin(np.linspace(0, 440 * 2 * np.pi * 5, sr * 5)).astype(np.float32)
    sf.write(config.VOICES_DIR / "myvoice.wav", tone, sr)
    (config.VOICES_DIR / "myvoice.txt").write_text("샘플 대본\n", encoding="utf-8")

    heard = {}  # 텍스트 → 받아쓰기 결과를 바꾸고 싶을 때

    cut = set()  # 여기 넣은 글자는 끝을 뚝 끊어서 만든다

    def fake_generate(text, voice_wav, voice_text, temperature, seed):
        n = int(sr * (0.5 + 0.05 * len(text)))
        a = (0.2 * np.sin(np.linspace(0, 300 * 2 * np.pi * n / sr, n))).astype(np.float32)
        if text not in cut:
            a[-int(0.15 * sr):] *= np.geomspace(1, 1e-4, int(0.15 * sr)).astype(np.float32)  # 사그라듦
        return np.concatenate([a, np.zeros(sr // 5, dtype=np.float32)])

    last = {}

    def fake_generate_wrap(text, *a):
        last["text"] = text
        return fake_generate(text, *a)

    def fake_transcribe(path):
        t = last["text"]
        return heard.get(t, t)

    monkeypatch.setattr(engine, "generate", fake_generate_wrap)
    monkeypatch.setattr(engine, "max_attempts", lambda: config.MAX_AUTO_ATTEMPTS)
    monkeypatch.setattr(engine, "transcribe", fake_transcribe)
    return {"heard": heard, "tmp": tmp_path, "cut": cut}


# ── 순수 함수 ──


def test_split_keeps_lines_and_sentences():
    chunks = split_script(SCRIPT)
    assert [c[1] for c in chunks] == [
        "오늘은 이 공허의 약탈자를, 게임이 아니라 생물학의 눈으로 한 번 뜯어보겠습니다.",
        "먼저 몸부터 보죠.",
        "몸을 빠르게 움직이는 뒤쪽 다리",
        "사마귀입니다.",
    ]
    assert [c[0] for c in chunks] == [0, 0, 0, 1]
    assert split_script(join_script(chunks)) == chunks


def test_split_long_sentence_at_comma():
    s = ("가" * 70 + ", ") + ("나" * 70) + "."
    chunks = split_script(s)
    assert len(chunks) == 2 and all(len(c[1]) <= 120 for c in chunks)


@pytest.mark.parametrize("script,heard,ok", [
    ("그리고 먹이를 찢어내는, 낫처럼 뻗은 앞 다리", "그리고 먹이를 뛰어내는 나처럼 뻗은 앞다리.", False),
    ("그리고 먹이를 찢어내는, 낫처럼 뻗은 앞 다리", "그리고 먹이를 찢어내는 나처럼 뻗은 앞다리.", True),
    ("오늘은 이 공허의 약탈자를", "오늘은 이 공어의 약탈자를", True),
    ("사마귀입니다.", "사마귀입니다", True),
    ("사마귀입니다.", "사막이입니다.", False),  # 짧은 줄은 더 엄격
    ("다시 놓고 보면, 자연스럽게 떠오르는 곤충이 있습니다.", "다시 놓고 오면 자연스럽게 떠오르는 곤충이 있습니다.", True),
    ("사마귀입니다.", "", False),
])
def test_check_score(script, heard, ok):
    assert check.score(script, heard)["ok"] is ok


def test_check_diff_marks_wrong_part():
    diff = check.score("찢어내는, 낫처럼 다리", "뛰어내는 나처럼 다리")["diff"]
    assert diff == [["찢", "뛰"], ["낫", "나"]]


def test_srt_format():
    srt = subtitles.make_srt([(0, 1.5, "가"), (61.25, 62, "나")])
    assert "00:00:00,000 --> 00:00:01,500" in srt and "00:01:01,250 --> 00:01:02,000" in srt


def test_subtitle_wrap_fits_width():
    font = subtitles._font(64)
    lines = subtitles.wrap("아주 " * 60, font, 1600)
    assert len(lines) > 1 and all(font.getlength(l) <= 1600 for l in lines)


# ── 저장 ──


def test_set_script_keeps_matching_lines(env):
    d = store.create_project("ep01", SCRIPT)
    first = d["lines"][0]
    first["speed"] = 1.1
    store.set_script(d, "새 첫 줄.\n" + store.script_of(d))
    assert d["lines"][1]["id"] == first["id"] and d["lines"][1]["speed"] == 1.1
    assert d["lines"][0]["text"] == "새 첫 줄."


def test_bad_project_id_is_not_found(env):
    with pytest.raises(store.NotFound):
        store.load("../../etc")


def test_voice_name_rules(env):
    with pytest.raises(store.BadRequest):
        store.voice_paths("../x")
    assert store.voice_paths("차분_1")[0].name == "차분_1.wav"


# ── 생성 흐름 (가짜 모델) ──


def run_queue(pid, lids, manual=False):
    for lid in lids:
        worker._run_line(pid, lid, manual)


def test_generate_check_and_retry(env):
    d = store.create_project("ep01", SCRIPT)
    bad = d["lines"][2]["text"]
    env["heard"][bad] = "전혀 다른 말이 들렸습니다"
    run_queue(d["id"], [l["id"] for l in d["lines"]])
    d = store.load(d["id"])
    views = [store.line_view(d, l) for l in d["lines"]]
    assert [v["check"] for v in views] == ["ok", "ok", "bad", "ok"]
    assert len(views[2]["takes"]) == config.MAX_AUTO_ATTEMPTS  # 틀리면 자동으로 끝까지 다시
    assert len(views[0]["takes"]) == 1  # 맞으면 한 번만

    # 무시하면 통과 처리
    d["lines"][2]["ignore_check"] = True
    assert store.line_view(d, d["lines"][2])["check"] == "ok"


def test_manual_redo_adds_take_and_not_regenerated_when_done(env):
    d = store.create_project("ep01", SCRIPT)
    lid = d["lines"][0]["id"]
    run_queue(d["id"], [lid])
    run_queue(d["id"], [lid])  # 자동: 이미 있으면 건너뜀
    assert len(store.load(d["id"])["lines"][0]["takes"]) == 1
    run_queue(d["id"], [lid], manual=True)
    d = store.load(d["id"])
    assert len(d["lines"][0]["takes"]) == 2 and d["lines"][0]["chosen"] == d["lines"][0]["takes"][-1]["id"]


def test_text_edit_invalidates_take(env):
    d = store.create_project("ep01", SCRIPT)
    lid = d["lines"][0]["id"]
    run_queue(d["id"], [lid])
    d = store.load(d["id"])
    d["lines"][0]["text"] = "바뀐 문장."
    assert store.chosen_take(d, d["lines"][0]) is None


# ── 내보내기 ──


def test_export_files_and_timing(env):
    d = store.create_project("ep01 카직스", SCRIPT)
    d["lines"][1]["speed"] = 1.25
    d["lines"][2]["gap"] = 2.0
    store.save(d)
    run_queue(d["id"], [l["id"] for l in d["lines"]])
    d = store.load(d["id"])
    with pytest.raises(exporter.NotReady):
        dd = store.load(d["id"])
        dd["lines"][0]["chosen"] = None
        exporter.build(dd)
    res = exporter.export(d)
    out = Path(res["local"])
    assert sorted(p.name for p in out.iterdir()) == ["lines", "narration.srt", "narration.wav", "subtitles_green.mp4"]
    assert len(list((out / "lines").glob("*.wav"))) == 4
    assert res["icloud"] and (Path(res["icloud"]) / "narration.wav").exists()

    wav = sf.info(str(out / "narration.wav"))
    assert wav.samplerate == 48000
    _, cues, _ = exporter.build(d)
    assert abs(wav.duration - cues[-1][1]) < 0.05
    # 줄 3 뒤에는 2초 쉼
    assert abs((cues[3][0] - cues[2][1]) - 2.0) < 0.01
    # 줄 2는 1.25배 빠르게
    t1 = next(t for t in d["lines"][1]["takes"] if t["id"] == d["lines"][1]["chosen"])
    assert abs((cues[1][1] - cues[1][0]) - t1["duration"] / 1.25) < 0.1

    import subprocess
    dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                str(out / "subtitles_green.mp4")], capture_output=True, text=True).stdout)
    assert abs(dur - wav.duration) < 0.1


# ── 서버 ──


@pytest.fixture
def client(env, monkeypatch):
    from fastapi.testclient import TestClient

    from server import app as appmod

    monkeypatch.setattr(worker, "enqueue", lambda pid, lids, manual=False: len(lids))
    return TestClient(appmod.app, base_url="http://127.0.0.1:7870")


def test_api_security_headers_and_host(client, monkeypatch):
    r = client.get("/api/projects")
    assert r.status_code == 200
    assert "default-src 'self'" in r.headers["content-security-policy"]
    assert r.headers["x-frame-options"] == "DENY"
    monkeypatch.delenv("STUDIO_TESTING")
    assert client.get("/api/projects", headers={"host": "evil.com"}).status_code == 403
    assert client.post("/api/projects", json={"title": "x"}, headers={"origin": "http://evil.com"}).status_code == 403
    assert client.post("/api/projects", json={"title": "x"}, headers={"origin": "http://127.0.0.1:7870"}).status_code == 200


def test_api_project_flow(client):
    p = client.post("/api/projects", json={"title": "ep01", "script": SCRIPT}).json()
    assert len(p["lines"]) == 4 and p["settings"]["voice"] == "myvoice"
    lid = p["lines"][0]["id"]
    r = client.patch(f"/api/projects/{p['id']}/lines/{lid}", json={"speed": 1.1, "gap": 0.5})
    assert r.json()["speed"] == 1.1 and r.json()["gap"] == 0.5
    assert client.patch(f"/api/projects/{p['id']}/lines/{lid}", json={"voice": "없는목소리"}).status_code == 400
    r = client.patch(f"/api/projects/{p['id']}", json={"settings": {"subtitle": {"font_size": 80, "position": "top"}}})
    assert r.json()["settings"]["subtitle"]["font_size"] == 80
    assert client.patch(f"/api/projects/{p['id']}", json={"settings": {"subtitle": {"color": "red"}}}).status_code == 400
    r = client.put(f"/api/projects/{p['id']}/script", json={"script": "하나.\n\n둘."})
    assert [l["text"] for l in r.json()["lines"]] == ["하나.", "둘."]
    assert client.get(f"/api/projects/{p['id']}/subtitle-preview.png").headers["content-type"] == "image/png"
    assert client.post(f"/api/projects/{p['id']}/export").status_code == 409  # 아직 안 만듦
    assert client.get("/api/projects/nope").status_code == 404
    assert client.get(f"/api/projects/{p['id']}/export/file", params={"path": "../project.json"}).status_code == 404


def test_api_voice_upload(client, env, monkeypatch):
    monkeypatch.setattr(engine, "transcribe", lambda path: "받아쓴 대본")
    src = env["tmp"] / "rec.wav"
    sf.write(src, 0.3 * np.sin(np.linspace(0, 2000, config.SR * 8)).astype(np.float32), config.SR)
    with open(src, "rb") as f:
        r = client.post("/api/voices", data={"name": "차분"}, files={"file": ("rec.wav", f, "audio/wav")})
    assert r.status_code == 200, r.text
    assert r.json()["text"] == "받아쓴 대본"
    names = [v["name"] for v in client.get("/api/voices").json()]
    assert names == ["myvoice", "차분"]
    assert client.patch("/api/voices/차분", json={"text": "고친 대본"}).json()["text"] == "고친 대본"
    client.post("/api/projects", json={"title": "ep", "script": "가."})
    assert client.delete("/api/voices/myvoice").status_code == 400  # 쓰는 에피소드 있음
    assert client.delete("/api/voices/차분").status_code == 200


def test_subtitle_balanced_wrap_evens_lines():
    font = subtitles._font(64)
    text = "오늘은 이 공허의 약탈자를, 게임이 아니라 생물학의 눈으로 한 번 뜯어보겠습니다."
    lines = subtitles.balanced_wrap(text, font, 1920 * 0.86)
    assert len(lines) == len(subtitles.wrap(text, font, 1920 * 0.86)) == 2
    widths = [font.getlength(l) for l in lines]
    assert min(widths) / max(widths) > 0.7


def test_trim_keeps_fading_tail():
    """끝음이 사그라드는 꼬리(-54dB 위)는 남기고, 그 뒤 무음만 자른다."""
    from server import audio

    sr = config.SR
    speech = 0.3 * np.ones(sr // 2, dtype=np.float32)
    fade = np.linspace(0.3, 0.003, int(0.2 * sr), dtype=np.float32)  # 0.2초 동안 사그라듦
    silence = np.zeros(sr, dtype=np.float32)
    out = audio.trim_silence(np.concatenate([silence, speech, fade, silence]))
    assert len(out) >= len(speech) + len(fade)  # 꼬리가 잘리지 않음
    assert len(out) < len(speech) + len(fade) + int(0.25 * sr)  # 뒤 무음은 잘림
    assert abs(out[-1]) < 1e-6  # 끝은 페이드아웃


def test_api_generate_all_redo(client, monkeypatch):
    calls = []
    monkeypatch.setattr(worker, "enqueue", lambda pid, lids, manual=False: calls.append((len(lids), manual)) or len(lids))
    p = client.post("/api/projects", json={"title": "ep", "script": SCRIPT}).json()
    assert client.post(f"/api/projects/{p['id']}/generate").json()["added"] == 4
    assert client.post(f"/api/projects/{p['id']}/generate?redo=true").json()["added"] == 4
    assert calls == [(4, False), (4, True)]  # 다시 뽑기는 이미 있는 줄도 새 테이크로


@pytest.mark.parametrize("text,sent", [
    ("먼저 몸부터 보죠.", "먼저 몸부터 보죠..."),
    ("몸을 빠르게 움직이는 뒤쪽 다리", "몸을 빠르게 움직이는 뒤쪽 다리..."),
    ("어디까지 진화할 수 있을까요?", "어디까지 진화할 수 있을까요?.."),
    ("그렇구나…", "그렇구나..."),
])
def test_tts_text_adds_tail(text, sent):
    """문장 끝이 끊기지 않게 모델에만 말줄임표를 붙인다 (맥 엔진)."""
    assert engine.tts_text(text, "ellipsis") == sent


@pytest.mark.parametrize("text,sent", [
    ("몸을 빠르게 움직이는 뒤쪽 다리", "몸을 빠르게 움직이는 뒤쪽 다리."),
    ("먼저 몸부터 보죠.", "먼저 몸부터 보죠."),
    ("정말 그럴까요?", "정말 그럴까요?"),
])
def test_tts_text_period_style(text, sent):
    """PyTorch 엔진은 마침표만 채운다."""
    assert engine.tts_text(text, "period") == sent


def test_max_frames_scales_with_text():
    assert engine.max_frames("가" * 10) < engine.max_frames("가" * 100) < 2048


def test_ending_decay_detects_cut():
    from server import audio

    sr = config.SR
    tone = (0.2 * np.sin(np.linspace(0, 300 * 2 * np.pi, sr))).astype(np.float32)
    faded = tone.copy()
    faded[-int(0.15 * sr):] *= np.geomspace(1, 1e-4, int(0.15 * sr)).astype(np.float32)
    silence = np.zeros(sr // 2, dtype=np.float32)
    assert audio.ending_decay_ms(np.concatenate([faded, silence])) >= audio.ENDING_MIN_MS
    assert audio.ending_decay_ms(np.concatenate([tone, silence])) < audio.ENDING_MIN_MS  # 뚝 끊김
    assert audio.ending_decay_ms(tone) < audio.ENDING_MIN_MS  # 소리 난 채로 파일이 끝남


def test_cut_ending_is_regenerated_and_flagged(env):
    d = store.create_project("ep01", SCRIPT)
    l = d["lines"][1]
    env["cut"].add(l["text"])
    run_queue(d["id"], [l["id"]])
    d = store.load(d["id"])
    v = store.line_view(d, d["lines"][1])
    assert len(v["takes"]) == config.MAX_AUTO_ATTEMPTS  # 끝이 끊기면 자동으로 다시
    assert v["check"] == "bad"
    t = v["takes"][-1]
    assert t["tail_ok"] is False and t["pron_ok"] is True  # 발음은 맞고 끝만 끊김


# ── 대본 불러오기 ──


def _docx_bytes(paras):
    import io

    import docx

    doc = docx.Document()
    for p in paras:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(lines):
    """글자가 든 아주 작은 PDF (ASCII). 줄이 화면 폭에서 꺾인 것처럼 만든다."""
    body = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({ln}) Tj 0 -16 Td" for ln in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(body)} >>\nstream\n{body}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offs = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF"
    return out.encode("latin-1")


def test_import_docx_keeps_lines_and_paragraphs():
    from server import importer

    data = _docx_bytes(["먼저 몸부터 보죠.", "몸을 빠르게 움직이는 뒤쪽 다리", "", "", "사마귀입니다.", "3"])
    text = importer.extract_text(data, "ep02 사마귀.docx")
    assert text == "먼저 몸부터 보죠.\n몸을 빠르게 움직이는 뒤쪽 다리\n\n사마귀입니다."  # 빈 단락 = 문단, 쪽 번호는 뺌
    assert importer.title_from("ep02_사마귀.docx") == "ep02 사마귀"


def test_import_pdf_rejoins_wrapped_lines():
    from server import importer

    data = _pdf_bytes(["This sentence is wrapped", "across two lines.", "Second one."])
    assert importer.extract_text(data, "a.pdf") == "This sentence is wrapped across two lines.\nSecond one."


def test_import_txt_cp949_and_bad_types():
    from server import importer

    assert importer.extract_text("가나다.\n\n\n라마.".encode("cp949"), "a.txt") == "가나다.\n\n라마."
    for name in ("a.doc", "a.hwp", "a.png"):
        with pytest.raises(importer.ImportError_):
            importer.extract_text(b"x", name)


def test_inbox_scan_creates_then_updates(env, monkeypatch):
    import os

    from server import importer

    monkeypatch.setattr(importer, "SETTINGS_FILE", env["tmp"] / "settings.json")
    queued = []
    monkeypatch.setattr(worker, "enqueue", lambda pid, lids, manual=False: queued.append((pid, len(lids))) or len(lids))
    inbox = env["tmp"] / "drive" / "TTS 대본"
    inbox.mkdir(parents=True)
    importer.save_settings({"folder": str(inbox), "auto_generate": True, "auto_export": False})
    f = inbox / "ep02 사마귀.docx"
    f.write_bytes(_docx_bytes(["먼저 몸부터 보죠.", "사마귀입니다."]))
    os.utime(f, (1_000_000_000, 1_000_000_000))  # '내려받는 중' 이 아니게 (5초 전보다 오래된 파일)
    assert worker.scan_inbox() == 1
    assert worker.scan_inbox() == 0  # 안 바뀌면 다시 안 불러옴
    [p] = store.list_projects()
    assert p["title"] == "ep02 사마귀" and p["lines"] == 2 and queued == [(p["id"], 2)]

    f.write_bytes(_docx_bytes(["먼저 몸부터 보죠.", "사마귀입니다.", "끝."]))
    os.utime(f, (1_000_000_100, 1_000_000_100))
    assert worker.scan_inbox() == 1
    [p2] = store.list_projects()  # 새로 만들지 않고 같은 에피소드의 대본을 바꾼다
    assert p2["id"] == p["id"] and p2["lines"] == 3


def test_api_import_project_and_parse(client, env, monkeypatch):
    from server import importer

    monkeypatch.setattr(importer, "SETTINGS_FILE", env["tmp"] / "settings.json")
    r = client.post("/api/import/parse", files={"file": ("ep03.docx", _docx_bytes(["하나.", "둘."]), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["text"] == "하나.\n둘." and r.json()["title"] == "ep03"
    assert client.post("/api/import/parse", files={"file": ("a.hwp", b"x", "application/octet-stream")}).status_code == 400
    r = client.post("/api/projects/import", data={"text": "가.\n나.", "title": "드라이브 문서", "generate": "false"})
    assert r.status_code == 200 and r.json()["title"] == "드라이브 문서" and len(r.json()["lines"]) == 2
    folder = env["tmp"] / "drive" / "TTS 대본"
    folder.parent.mkdir()
    r = client.put("/api/inbox", json={"folder": str(folder), "auto_generate": False, "auto_export": False})
    assert r.status_code == 200 and folder.is_dir()  # 없으면 만들어 준다
    assert client.put("/api/inbox", json={"folder": "상대/경로"}).status_code == 400


def _docx_table_bytes(parts):
    """제작용 대본처럼: 앞에 제작 메모 단락, PART 마다 표 (타임코드 | 장면 | 나레이션 (녹음용) | 화면)."""
    import io

    import docx

    doc = docx.Document()
    doc.add_paragraph("제작 원칙: 이 줄은 대본이 아니다.")
    for rows in parts:
        t = doc.add_table(rows=1, cols=4)
        for c, h in zip(t.rows[0].cells, ["타임코드", "장면 / 의도", "나레이션 (녹음용)", "화면 / 자료"]):
            c.text = h
        for tc, narr in rows:
            r = t.add_row().cells
            r[0].text, r[1].text, r[3].text = tc, "장면", "카메라 설명"
            r[2].text = narr[0]
            for extra in narr[1:]:
                r[2].add_paragraph(extra)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_import_docx_production_table_takes_narration_column():
    from server import importer

    data = _docx_table_bytes([
        [("0:00", ["여기, 협곡을 누비는 괴물이 있습니다."]), ("0:25", ["하지만 단어 하나는 떠오릅니다.벌레."])],
        [("0:47", ["먼저 몸부터 보죠.", "몸을 빠르게 움직이는 뒤쪽 다리."])],
    ])
    assert importer.extract_text(data, "script.docx") == (
        "여기, 협곡을 누비는 괴물이 있습니다.\n하지만 단어 하나는 떠오릅니다.\n벌레.\n\n"
        "먼저 몸부터 보죠.\n몸을 빠르게 움직이는 뒤쪽 다리."
    )


def test_import_markdown_table_from_drive_text():
    from server import importer

    md = """조성주의 바이올로지 - 제작용 대본

| 예상 러닝타임 | 9분 | 콘텐츠 형식 | 영상 에세이 |
| :-: | :-: | :-: | :-: |
| 핵심 구조 | 훅 | 핵심 타깃 | 10\\~30대 |

PART 1. 협곡에 사는 괴물

| 타임코드 | 장면 / 의도 | 나레이션 (녹음용) | 화면 / 자료 |
| :-: | :-: | :-: | :-: |
| 0:00–0:05 | 콜드 오픈 | 여기, 협곡을 누비는 괴물이 있습니다. | 카직스 이동 |
| 0:25–0:32 | 훅 | 하지만 이 모든 행동을 보고도, 꽤 잘 어울리는 단어 하나는 떠오릅니다.벌레. | \\`벌레.\\` |

PART 2. 메뚜기인가

| 타임코드 | 장면 | 나레이션 (녹음용) | 화면 |
| :-: | :-: | :-: | :-: |
| 1:38–1:52 | 반론 | 그런데 롤을 해본 사람이라면 조금 이상할 겁니다.카직스는 보통… 메뚜기라고 불리거든요. | 점프 |
"""
    assert importer.extract_text(md.encode(), "drive.md") == (
        "여기, 협곡을 누비는 괴물이 있습니다.\n"
        "하지만 이 모든 행동을 보고도, 꽤 잘 어울리는 단어 하나는 떠오릅니다.\n벌레.\n\n"
        "그런데 롤을 해본 사람이라면 조금 이상할 겁니다.\n카직스는 보통… 메뚜기라고 불리거든요."
    )
    # 표가 없는 평범한 텍스트는 그대로
    assert importer.extract_text("가.\n나.".encode(), "a.md") == "가.\n나."


def test_import_md_flattened_cell_lines():
    """드라이브 글자에서는 칸 안 줄바꿈이 사라져 '빌리고,때로는' 처럼 붙어 온다 → 줄을 되살린다."""
    from server import importer

    md = """| 타임코드 | 나레이션 (녹음용) |
| :-: | :-: |
| 9:55 | 생명은 모든 것을 새로 발명하지 않았습니다.때로는 빌리고,때로는 합치고,그리고 다시 썼습니다. |
| 5:28 | 물론 먹었다고 \\`+1 진화 포인트\\`가 뜨지는 않습니다. |
"""
    assert importer.extract_text(md.encode(), "a.md") == (
        "생명은 모든 것을 새로 발명하지 않았습니다.\n때로는 빌리고,\n때로는 합치고,\n그리고 다시 썼습니다.\n"
        "물론 먹었다고 +1 진화 포인트가 뜨지는 않습니다."
    )


def test_split_keeps_ellipsis_inside_line():
    assert [c[1] for c in split_script("카직스는 보통… 메뚜기라고 불리거든요. 아니면… 바퀴벌레일까요?")] == [
        "카직스는 보통… 메뚜기라고 불리거든요.",
        "아니면… 바퀴벌레일까요?",
    ]
